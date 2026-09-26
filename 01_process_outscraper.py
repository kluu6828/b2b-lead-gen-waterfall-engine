import pandas as pd
import numpy as np
import argparse
import os
import re
from urllib.parse import urlparse, unquote
from openpyxl.utils import get_column_letter
from openpyxl import Workbook, load_workbook

MASTER_SUMMARY_FILENAME = "master_market_analysis_summary.xlsx"
MASTER_SHEET_NAME = "Master Summary"

SUMMARY_METRICS = [
    "Total Fitness Clubs / Gyms (Filtered & Unique Addresses)",
    "Independent Standalone (1 Location)",
    "National / Franchise Locations",
    "Independent Multi-Unit (2-4 Locations)",
    "Regional Accounts (5+ Locations)",
]

# ==========================================
# KNOWN NATIONAL / FRANCHISE GYM BRANDS
# ==========================================
KNOWN_NATIONAL_BRANDS_MAP = {
    "crossfit": "CrossFit",
    "ymca": "YMCA",
    "ywca": "YWCA",
    "planet fitness": "Planet Fitness",
    "anytime fitness": "Anytime Fitness",
    "crunch": "Crunch Fitness",
    "crunch fitness": "Crunch Fitness",
    "f45": "F45 Training",
    "f45 training": "F45 Training",
    "jazzercise": "Jazzercise",
    "life time": "Life Time",
    "lifetime fitness": "Life Time",
    "stretchmed": "StretchMed",
    "equinox": "Equinox",
    "la fitness": "LA Fitness",
    "orangetheory": "OrangeTheory Fitness",
    "orangetheory fitness": "OrangeTheory Fitness",
    "gold's gym": "Gold's Gym",
    "golds gym": "Gold's Gym",
    "snap fitness": "Snap Fitness",
    "24 hour fitness": "24 Hour Fitness",
    "workout anytime": "Workout Anytime"
}

REGIONAL_MIN_LOCATIONS = 5
BRAND_UMBRELLA_GROUP_KEYS = {"crossfit", "ymca"}  # Force all locations into one group regardless of affiliate name/domain
SOCIAL_MEDIA_DOMAINS = {"instagram.com", "facebook.com", "fb.com"}

def matches_known_national_brand(text):
    """Return (brand_key, canonical_name) if text matches a known national franchise, else None."""
    if not isinstance(text, str) or not text.strip():
        return None
    text_lower = text.lower()
    for brand_key, canonical_name in KNOWN_NATIONAL_BRANDS_MAP.items():
        if brand_key in text_lower:
            return brand_key, canonical_name
    return None

def sort_by_name_az(df, column):
    """Sort dataframe A-Z by column, case-insensitive."""
    return df.sort_values(by=column, key=lambda s: s.astype(str).str.lower()).reset_index(drop=True)

def extract_root_domain(url_str):
    """Extracts root domain (e.g., locations.crunch.com/tx/austin -> crunch.com)"""
    if not isinstance(url_str, str) or not url_str.strip():
        return ''
    url = url_str.strip().lower()
    if not url.startswith(('http://', 'https://')):
        url = 'http://' + url
    try:
        netloc = urlparse(url).netloc
        netloc = re.sub(r'^www\.', '', netloc)
        parts = netloc.split('.')
        if len(parts) >= 2:
            if len(parts) > 2 and parts[-2] in ['co', 'com', 'org', 'gov', 'edu', 'net']:
                return '.'.join(parts[-3:])
            return '.'.join(parts[-2:])
        return netloc
    except Exception:
        return ''

def extract_website_group_key(url_str):
    """
    Build a website grouping key for regional/independent chains.
    - Custom domains: root domain only (strip everything after first slash)
    - Instagram/Facebook: domain + page username (strip everything after username)
    """
    if not isinstance(url_str, str) or not url_str.strip():
        return ''

    url = url_str.strip().lower()
    if not url.startswith(('http://', 'https://')):
        url = 'http://' + url

    try:
        parsed = urlparse(url)
        root_domain = extract_root_domain(url)
        if not root_domain:
            return ''

        path_parts = [part for part in parsed.path.strip('/').split('/') if part]

        if root_domain in SOCIAL_MEDIA_DOMAINS:
            if not path_parts:
                return ''
            username = unquote(path_parts[0]).split('?')[0].strip('/')
            if not username:
                return ''
            return f"{root_domain}/{username}"

        return root_domain
    except Exception:
        return ''

def format_trackable_website(url_str):
    """Format website for display: root domain, or instagram.com/user / facebook.com/user."""
    key = extract_website_group_key(url_str)
    if key:
        if key.startswith('fb.com/'):
            return key.replace('fb.com/', 'facebook.com/', 1)
        return key
    return extract_root_domain(url_str)

def derive_group_chain_name(names):
    """Derive a shared brand name from location-specific names (e.g. drop city suffix)."""
    names = [str(name).strip() for name in names if isinstance(name, str) and str(name).strip()]
    if not names:
        return ''
    if len(names) == 1:
        return clean_chain_name(names[0])

    word_lists = [name.split() for name in names]
    common_words = []
    for i in range(min(len(words) for words in word_lists)):
        tokens = {words[i].lower() for words in word_lists}
        if len(tokens) == 1:
            common_words.append(word_lists[0][i])
        else:
            break

    if common_words:
        return clean_chain_name(' '.join(common_words))
    return clean_chain_name(names[0])

def clean_chain_name(raw_name):
    """Normalizes gym names by matching canonical franchises or stripping city/town suffixes."""
    if not isinstance(raw_name, str) or not raw_name.strip():
        return ''
    
    name_lower = raw_name.strip().lower()

    # Match against Known National Franchises
    for brand_key, canonical_name in KNOWN_NATIONAL_BRANDS_MAP.items():
        if brand_key in name_lower:
            return canonical_name

    # Clean generic location suffixes (e.g. "Texas Fitness - Austin" -> "Texas Fitness")
    cleaned = raw_name.strip()
    cleaned = re.sub(r'\s*[-|–—:,]\s*.*$', '', cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r'\s+\([^)]*\)', '', cleaned)
    
    return cleaned.strip().title()


def derive_state_prefix(file_path):
    basename = os.path.basename(file_path).rsplit(".", 1)[0]
    return basename.replace("_market_analysis", "").replace("_fit_scored", "")


def derive_state_slug(state_prefix):
    match = re.match(r"^([a-z]{2})-", state_prefix, re.I)
    if match:
        return match.group(1).lower()
    return state_prefix.split("-")[0].lower()


def resolve_state_output_dir(file_path, output_dir="cleaned_output"):
    state_prefix = derive_state_prefix(file_path)
    state_slug = derive_state_slug(state_prefix)
    state_output_dir = os.path.join(output_dir, state_slug)
    os.makedirs(state_output_dir, exist_ok=True)
    return state_output_dir, state_prefix, state_slug


def append_sheet_total_row(ws, label="TOTAL", label_col=1, sum_col_name="total_locations"):
    """Append a TOTAL row that sums total_locations via an Excel formula."""
    headers = {cell.value: cell.column for cell in ws[1]}
    sum_col = headers.get(sum_col_name)
    if not sum_col:
        raise ValueError(f"Column '{sum_col_name}' not found on sheet '{ws.title}'")

    last_data_row = ws.max_row
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=label_col, value=label)
    col_letter = get_column_letter(sum_col)
    ws.cell(
        row=total_row,
        column=sum_col,
        value=f"=SUM({col_letter}2:{col_letter}{last_data_row})",
    )
    return total_row, col_letter


def write_summary_page(wb, sheet_totals):
    """
    Build Summary Page with live formulas referencing account-tab TOTAL rows.

    Row order:
      Total fitness clubs | Independent standalone | National/franchise |
      Independent multi-unit | Regional (5+)
    """
    if "Summary Page" in wb.sheetnames:
        del wb["Summary Page"]

    ws = wb.create_sheet("Summary Page", 0)
    ws.append(["Metric", "Total Clubs", "%"])

    def cell_ref(info):
        return f"'{info['sheet']}'!{info['col_letter']}{info['total_row']}"

    rows = []
    segment_keys = ["standalone", "national", "multi", "regional"]
    for idx, metric in enumerate(SUMMARY_METRICS):
        if idx == 0:
            total_formula = "=B3+B4+B5+B6"
            pct_formula = "=B2/$B$2"
        else:
            total_formula = f"={cell_ref(sheet_totals[segment_keys[idx - 1]])}"
            pct_formula = f"=B{idx + 2}/$B$2"
        rows.append((metric, total_formula, pct_formula))

    for metric, total_formula, pct_formula in rows:
        ws.append([metric, total_formula, pct_formula])

    for row_idx in range(2, 7):
        ws.cell(row=row_idx, column=3).number_format = "0.0%"


def discover_state_market_analysis_files(output_dir="cleaned_output"):
    """Find all per-state market analysis workbooks under cleaned_output/{state}/."""
    if not os.path.isdir(output_dir):
        return []

    discovered = []
    for entry in os.listdir(output_dir):
        state_dir = os.path.join(output_dir, entry)
        if not os.path.isdir(state_dir) or entry == "trial":
            continue

        for filename in os.listdir(state_dir):
            if not filename.endswith("_market_analysis.xlsx") or filename.startswith("~$"):
                continue
            discovered.append({
                "state_slug": entry.lower(),
                "path": os.path.join(state_dir, filename),
            })
            break

    discovered.sort(key=lambda item: item["state_slug"])
    return discovered


def _sheet_location_total(ws, sum_col_name="total_locations"):
    """Sum total_locations on a tab, excluding the TOTAL summary row."""
    headers = {cell.value: cell.column for cell in ws[1]}
    sum_col = headers.get(sum_col_name)
    if not sum_col:
        return 0

    total = 0
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "TOTAL":
            continue
        val = ws.cell(row=row, column=sum_col).value
        if isinstance(val, (int, float)) and not isinstance(val, bool):
            total += int(val)
    return total


def read_state_summary_metrics(state_workbook_path):
    """
    Read the five Summary Page metrics from a state workbook by summing
    total_locations on each account tab (works without Excel formula cache).
    """
    wb = load_workbook(state_workbook_path, data_only=True)
    required_sheets = [
        "Independent Standalone",
        "National Accounts",
        "Independent Multi-Units",
        "Regional Accounts",
    ]
    missing = [name for name in required_sheets if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"Missing sheet(s): {', '.join(missing)}")

    standalone = _sheet_location_total(wb["Independent Standalone"])
    national = _sheet_location_total(wb["National Accounts"])
    multi = _sheet_location_total(wb["Independent Multi-Units"])
    regional = _sheet_location_total(wb["Regional Accounts"])
    total = standalone + national + multi + regional
    return [total, standalone, national, multi, regional]


def write_master_summary_wide(ws, state_files):
    """
    Wide layout: metrics listed once in column A; each state gets Total Clubs + %
    columns to the right with the state code as a merged header above them.
    """
    ws.cell(row=2, column=1, value="Metric")

    col_idx = 0
    for state_info in state_files:
        col_total = 2 + col_idx * 2
        col_pct = col_total + 1
        state_slug = state_info["state_slug"].upper()
        col_total_letter = get_column_letter(col_total)

        try:
            metrics = read_state_summary_metrics(state_info["path"])
        except (ValueError, KeyError, OSError) as exc:
            print(f"⚠️  Skipping {state_slug} in master summary: {exc}")
            continue

        ws.cell(row=1, column=col_total, value=state_slug)
        ws.merge_cells(
            start_row=1,
            start_column=col_total,
            end_row=1,
            end_column=col_pct,
        )
        ws.cell(row=2, column=col_total, value="Total Clubs")
        ws.cell(row=2, column=col_pct, value="%")

        for offset, metric in enumerate(SUMMARY_METRICS):
            row_idx = 3 + offset
            if col_idx == 0:
                ws.cell(row=row_idx, column=1, value=metric)
            ws.cell(row=row_idx, column=col_total, value=metrics[offset])
            ws.cell(
                row=row_idx,
                column=col_pct,
                value=f"={col_total_letter}{row_idx}/${col_total_letter}$3",
            )
            ws.cell(row=row_idx, column=col_pct).number_format = "0.0%"

        col_idx += 1


def update_master_market_summary(output_dir="cleaned_output"):
    """
    Build or refresh the master workbook that mirrors each state's Summary Page.
    Scans cleaned_output for all state market analysis files; adds new states and
    refreshes formulas for existing ones.
    """
    state_files = discover_state_market_analysis_files(output_dir)
    master_path = os.path.join(output_dir, MASTER_SUMMARY_FILENAME)

    wb = Workbook()
    ws = wb.active
    ws.title = MASTER_SHEET_NAME

    if not state_files:
        ws.append(["No state market analysis workbooks found yet."])
        ws.append(["Run Step 1 for a state to populate this file."])
        wb.save(master_path)
        print(f"📊 Master summary initialized: {master_path}")
        return master_path

    write_master_summary_wide(ws, state_files)

    wb.save(master_path)
    state_list = ", ".join(s["state_slug"].upper() for s in state_files)
    print(f"📊 Master summary updated ({len(state_files)} states: {state_list})")
    print(f"   {master_path}")
    return master_path


def clean_and_segment_outscraper(file_path, output_dir="cleaned_output"):
    state_output_dir, state_prefix, state_slug = resolve_state_output_dir(file_path, output_dir)
    print(f"🗂️  State output folder: {state_output_dir}/")

    # -------------------------------------------------------------
    # STEP 1: GET RAW DATA
    # -------------------------------------------------------------
    print(f"📂 STEP 1: Loading raw Outscraper dataset: {file_path}")
    df = pd.read_csv(file_path)
    raw_total = len(df)

    category_col = next((c for c in ['category', 'type', 'subtypes'] if c in df.columns), None)
    address_col = next((c for c in ['full_address', 'street', 'address'] if c in df.columns), None)
    name_col = next((c for c in ['name', 'title'] if c in df.columns), None)
    site_col = next((c for c in ['site', 'domain', 'website'] if c in df.columns), None)
    phone_col = next((c for c in ['phone', 'phone_number'] if c in df.columns), None)

    if not address_col or not name_col:
        raise ValueError("❌ Error: Dataset must contain at least a business name and street address column.")

    # -------------------------------------------------------------
    # STEP 2: STRICT FILTER FOR 'FITNESS CENTER' OR 'GYM' ONLY
    # -------------------------------------------------------------
    if category_col:
        allowed_categories = {'fitness center', 'gym'}
        df = df[
            df[category_col].fillna('').astype(str).str.strip().str.lower().isin(allowed_categories)
        ].copy()
        filtered_total = len(df)
        print(f"🎯 STEP 2: Category Filter Applied -> Kept {filtered_total} exact 'fitness center' / 'gym' records (Discarded {raw_total - filtered_total} non-gym rows).")
    else:
        print("⚠️ Warning: No 'category' column detected. Proceeding with all records.")
        filtered_total = len(df)

    # Export filtered records before any deduplication or grouping
    filtered_output_path = os.path.join(state_output_dir, f"{state_prefix}_filtered_gyms_pre_dedup.xlsx")
    df = sort_by_name_az(df, name_col)
    df.to_excel(filtered_output_path, index=False)
    print(f"📋 Saved pre-dedup filtered spreadsheet ({filtered_total} rows): {filtered_output_path}")

    # -------------------------------------------------------------
    # STEP 3: DEDUPLICATE BY PHYSICAL STREET ADDRESS
    # -------------------------------------------------------------
    df['clean_address'] = df[address_col].fillna('').astype(str).str.lower().str.strip()
    df_unique = df.drop_duplicates(subset=['clean_address']).copy()
    total_unique_locations = len(df_unique)
    removed_duplicates = filtered_total - total_unique_locations

    print(f"🧹 STEP 3: Address Deduplication -> Removed {removed_duplicates} duplicate physical addresses (Unique total: {total_unique_locations}).")

    # -------------------------------------------------------------
    # STEP 4: GROUPING LOGIC
    # - CrossFit / YMCA: umbrella group (ignore per-location names & domains)
    # - Known national brands: group by brand name (e.g. LA Fitness)
    # - Custom-domain websites: group by root domain (ignore path after first slash)
    # - Instagram / Facebook: group by domain + page username
    # - Fallback: group by cleaned business name
    # -------------------------------------------------------------
    df_unique['canonical_chain_name'] = df_unique[name_col].apply(clean_chain_name)
    
    if site_col:
        df_unique['root_domain'] = df_unique[site_col].apply(extract_root_domain)
    else:
        df_unique['root_domain'] = ''

    def assign_grouping_and_website(row):
        chain_lower = str(row['canonical_chain_name']).lower()
        brand_match = matches_known_national_brand(chain_lower)

        if brand_match:
            brand_key, canonical_name = brand_match
            if brand_key in BRAND_UMBRELLA_GROUP_KEYS:
                return pd.Series([brand_key, canonical_name, "None"])
            return pd.Series([brand_key, canonical_name, row['root_domain']])

        website_url = row[site_col] if site_col else ''
        website_key = extract_website_group_key(website_url)
        if website_key:
            return pd.Series([website_key, row['canonical_chain_name'], website_key])

        group_key = chain_lower if chain_lower else str(row[name_col]).lower().strip()
        return pd.Series([group_key, row['canonical_chain_name'], row['root_domain']])

    df_unique[['group_key', 'canonical_chain_name', 'final_website']] = df_unique.apply(
        assign_grouping_and_website, axis=1
    )

    # Calculate cumulative location counts per group across the state dataset
    chain_counts = df_unique['group_key'].value_counts()
    df_unique['total_locations'] = df_unique['group_key'].map(chain_counts)

    # Account Categorization Rules
    def classify_record(row):
        name = str(row['canonical_chain_name']).lower()
        loc_count = row['total_locations']

        if matches_known_national_brand(name):
            return "National/Franchise"
        elif loc_count >= REGIONAL_MIN_LOCATIONS:
            return "Regional"
        else:
            return "Independent"

    df_unique['account_category'] = df_unique.apply(classify_record, axis=1)

    # -------------------------------------------------------------
    # BUILD DELIVERABLE TABS
    # -------------------------------------------------------------
    nat_locs = len(df_unique[df_unique['account_category'] == "National/Franchise"])
    reg_locs = len(df_unique[df_unique['account_category'] == "Regional"])
    ind_locs = len(df_unique[df_unique['account_category'] == "Independent"])

    independent_df = df_unique[df_unique['account_category'] == "Independent"]
    ind_standalone_locs = len(independent_df[independent_df['total_locations'] == 1])
    ind_multi_locs = len(independent_df[independent_df['total_locations'].between(2, 4)])

    # Tab 2: National Accounts (Grouped)
    national_summary = (
        df_unique[df_unique['account_category'] == "National/Franchise"]
        .groupby('group_key')
        .agg(
            chain_name=('canonical_chain_name', 'first'),
            website=('final_website', 'first'),
            total_locations=('total_locations', 'first')
        )
        .reset_index(drop=True)
        .pipe(sort_by_name_az, 'chain_name')
    )

    # Tab 3: Regional Accounts (Grouped)
    def summarize_regional_group(group):
        return pd.Series({
            'chain_name': derive_group_chain_name(group[name_col].tolist()),
            'website': group['final_website'].iloc[0],
            'total_locations': group['total_locations'].iloc[0],
        })

    regional_summary = (
        df_unique[df_unique['account_category'] == "Regional"]
        .groupby('group_key')
        .apply(summarize_regional_group, include_groups=False)
        .reset_index(drop=True)
        .pipe(sort_by_name_az, 'chain_name')
    )

    # Tab 4: Independent Accounts (Itemized individual venues)
    indie_cols = [c for c in [name_col, site_col, address_col, phone_col] if c is not None]
    independent_summary = independent_df[indie_cols].copy()
    independent_summary.rename(
        columns={
            name_col: 'chain_name',
            site_col: 'website',
            address_col: 'street_address'
        },
        inplace=True
    )
    if site_col:
        independent_summary['website'] = independent_summary['website'].apply(format_trackable_website)
    independent_summary = sort_by_name_az(independent_summary, 'chain_name')

    # Tab 5: Independent Standalone (1 unique location per group)
    def summarize_independent_group(group):
        return pd.Series({
            'chain_name': derive_group_chain_name(group[name_col].tolist()),
            'website': format_trackable_website(group['final_website'].iloc[0]),
            'total_locations': group['total_locations'].iloc[0],
        })

    independent_standalone = (
        independent_df[independent_df['total_locations'] == 1][indie_cols]
        .copy()
        .rename(
            columns={
                name_col: 'chain_name',
                site_col: 'website',
                address_col: 'street_address'
            }
        )
    )
    if site_col:
        independent_standalone['website'] = independent_standalone['website'].apply(format_trackable_website)
    independent_standalone['total_locations'] = 1
    independent_standalone = sort_by_name_az(independent_standalone, 'chain_name')

    # Tab 6: Independent Multi-Units (2-4 locations, grouped)
    independent_multi_unit = (
        independent_df[independent_df['total_locations'].between(2, 4)]
        .groupby('group_key')
        .apply(summarize_independent_group, include_groups=False)
        .reset_index(drop=True)
        .pipe(sort_by_name_az, 'chain_name')
    )

    # -------------------------------------------------------------
    # EXPORT TO SINGLE EXCEL WORKBOOK (6 TABS)
    # -------------------------------------------------------------
    excel_output_path = os.path.join(state_output_dir, f"{state_prefix}_market_analysis.xlsx")

    with pd.ExcelWriter(excel_output_path, engine='openpyxl') as writer:
        national_summary.to_excel(writer, sheet_name='National Accounts', index=False)
        regional_summary.to_excel(writer, sheet_name='Regional Accounts', index=False)
        independent_summary.to_excel(writer, sheet_name='Independent Accounts', index=False)
        independent_standalone.to_excel(writer, sheet_name='Independent Standalone', index=False)
        independent_multi_unit.to_excel(writer, sheet_name='Independent Multi-Units', index=False)

        wb = writer.book
        sheet_totals = {}
        for key, sheet_name in [
            ("national", "National Accounts"),
            ("regional", "Regional Accounts"),
            ("multi", "Independent Multi-Units"),
            ("standalone", "Independent Standalone"),
        ]:
            total_row, col_letter = append_sheet_total_row(wb[sheet_name])
            sheet_totals[key] = {
                "sheet": sheet_name,
                "total_row": total_row,
                "col_letter": col_letter,
            }
        write_summary_page(wb, sheet_totals)

    summary_preview = pd.DataFrame({
        "Metric": SUMMARY_METRICS,
        "Total Clubs": [
            total_unique_locations,
            ind_standalone_locs,
            nat_locs,
            ind_multi_locs,
            reg_locs,
        ],
        "%": [
            "100.0%",
            f"{(ind_standalone_locs / total_unique_locations * 100):.1f}%" if total_unique_locations else "0%",
            f"{(nat_locs / total_unique_locations * 100):.1f}%" if total_unique_locations else "0%",
            f"{(ind_multi_locs / total_unique_locations * 100):.1f}%" if total_unique_locations else "0%",
            f"{(reg_locs / total_unique_locations * 100):.1f}%" if total_unique_locations else "0%",
        ],
    })

    update_master_market_summary(output_dir)

    print("\n✅ Processing Complete! Output saved to:")
    print(f"  📄 Excel File: {excel_output_path}\n")
    print("Summary preview (Summary Page uses live Excel formulas):")
    print(summary_preview.to_string(index=False))

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Clean and Segment Outscraper Gym Data")
    parser.add_argument("--file", type=str, help="Path to raw Outscraper CSV")
    parser.add_argument(
        "--rebuild-master",
        action="store_true",
        help="Rebuild master market summary from all existing state workbooks",
    )
    args = parser.parse_args()

    if args.rebuild_master and args.file:
        parser.error("Use either --file or --rebuild-master, not both.")
    if args.rebuild_master:
        update_master_market_summary()
    elif args.file:
        clean_and_segment_outscraper(args.file)
    else:
        parser.error("Provide --file to process a state, or --rebuild-master to refresh the master summary.")